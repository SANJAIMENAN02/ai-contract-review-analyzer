"""
document_processor.py

Validates uploads and extracts raw text from contract files (PDF / DOCX / TXT).

Each failure mode has its own exception subclass carrying the HTTP status the
API should answer with, so the router does not need to guess.
"""
from __future__ import annotations

import os
import re
import zipfile

import docx  # python-docx
import fitz  # PyMuPDF

from app.utils import config


class DocumentProcessingError(Exception):
    """Base class: a document cannot be read or contains no usable text."""

    status_code = 422


class UnsupportedFileTypeError(DocumentProcessingError):
    status_code = 400


class EmptyDocumentError(DocumentProcessingError):
    status_code = 400


class FileTooLargeError(DocumentProcessingError):
    status_code = 413


class TextExtractionError(DocumentProcessingError):
    status_code = 422


# Kept for backwards compatibility with earlier code / tests.
ALLOWED_EXTENSIONS = set(config.ALLOWED_EXTENSIONS)
MIN_TEXT_CHARS = 50


def get_extension(filename: str) -> str:
    return os.path.splitext(os.path.basename(filename or ""))[1].lower()


def validate_file(filename: str, file_size: int) -> None:
    """Validate extension and size before any processing happens."""
    ext = get_extension(filename)
    if ext not in ALLOWED_EXTENSIONS:
        raise UnsupportedFileTypeError(
            f"Invalid file type '{ext or 'unknown'}'. Supported formats: PDF, DOCX, TXT."
        )
    if file_size <= 0:
        raise EmptyDocumentError("The uploaded file is empty.")
    if file_size > config.max_file_size_bytes():
        raise FileTooLargeError(
            f"File too large ({file_size / (1024 * 1024):.2f} MB). "
            f"Maximum allowed size is {config.max_file_size_mb()} MB."
        )


def validate_content_signature(filename: str, head: bytes) -> None:
    """
    Check the file's leading bytes match its claimed extension. The browser
    supplied MIME type / filename are never trusted on their own.
    """
    ext = get_extension(filename)
    if ext == ".pdf" and not head.lstrip().startswith(b"%PDF"):
        raise UnsupportedFileTypeError("The file content does not look like a valid PDF.")
    if ext == ".docx" and not head.startswith(b"PK\x03\x04"):
        raise UnsupportedFileTypeError("The file content does not look like a valid DOCX document.")
    if ext == ".txt" and b"\x00" in head:
        raise UnsupportedFileTypeError("The file content does not look like plain text.")


def make_safe_filename(contract_id: str, filename: str) -> str:
    """Internal storage name: never derived from the uploaded name except the extension."""
    return f"{contract_id}{get_extension(filename)}"


def extract_text_from_pdf(file_path: str) -> str:
    try:
        chunks = []
        with fitz.open(file_path) as pdf_doc:
            if pdf_doc.page_count == 0:
                raise TextExtractionError("The PDF file has no pages.")
            for page in pdf_doc:
                chunks.append(page.get_text("text"))
        return "\n".join(chunks)
    except DocumentProcessingError:
        raise
    except Exception as exc:  # noqa: BLE001
        raise TextExtractionError(
            "Unable to extract text from this PDF. It may be corrupted, "
            "password-protected, or a scanned image without embedded text."
        ) from exc


def extract_text_from_docx(file_path: str) -> str:
    try:
        document = docx.Document(file_path)
        parts = [p.text for p in document.paragraphs if p.text.strip()]
        # Tables often hold payment schedules, so include them.
        for table in document.tables:
            for row in table.rows:
                row_text = " | ".join(c.text.strip() for c in row.cells if c.text.strip())
                if row_text:
                    parts.append(row_text)
        return "\n".join(parts)
    except (zipfile.BadZipFile, KeyError, ValueError, OSError) as exc:
        raise TextExtractionError(
            "Unable to extract text from this DOCX file. It may be corrupted "
            "or not a valid Word document."
        ) from exc
    except Exception as exc:  # noqa: BLE001
        raise TextExtractionError("Unable to extract text from this DOCX file.") from exc


def extract_text_from_txt(file_path: str) -> str:
    with open(file_path, "rb") as f:
        raw = f.read()
    for encoding in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise TextExtractionError("Unable to decode this text file.")


def extract_text(file_path: str, filename: str) -> str:
    """Dispatch on extension, then clean the result."""
    ext = get_extension(filename)
    if ext == ".pdf":
        raw_text = extract_text_from_pdf(file_path)
    elif ext == ".docx":
        raw_text = extract_text_from_docx(file_path)
    elif ext == ".txt":
        raw_text = extract_text_from_txt(file_path)
    else:
        raise UnsupportedFileTypeError(f"Invalid file type '{ext}'.")

    cleaned = clean_text(raw_text)
    if not cleaned.strip():
        raise EmptyDocumentError("The document is empty.")
    if len(cleaned.strip()) < MIN_TEXT_CHARS:
        raise TextExtractionError(
            "Unable to extract enough readable text. The document may be a scanned "
            "image, nearly empty, or contain only non-text content."
        )
    return cleaned


def clean_text(text: str) -> str:
    """Normalize whitespace while preserving paragraph structure."""
    if not text:
        return ""
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\x00", "")
    text = re.sub(r"\n{3,}", "\n\n", text)
    text = re.sub(r"[ \t]{2,}", " ", text)
    text = "\n".join(line.strip() for line in text.split("\n"))
    return text.strip()
